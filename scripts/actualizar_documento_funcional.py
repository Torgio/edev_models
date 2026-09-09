"""Actualiza el Word funcional existente sin perder su diseño ni sus figuras."""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCX = Path("output/documentos/Documento_funcional_Pulso_Energia.docx")


def replace_text(paragraph, text: str, bold_prefix: str | None = None) -> None:
    paragraph.clear()
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def insert_like(target, template, text: str):
    paragraph = target.insert_paragraph_before()
    if template._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(template._p.pPr))
    replace_text(paragraph, text)
    return paragraph


def insert_heading(target, text: str, level: int):
    return target.insert_paragraph_before(text, style=f"Heading {level}")


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def set_table_cell_text(cell, text: str) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(31, 74, 68)


def paragraph_by_text(doc, exact: str):
    for paragraph in doc.paragraphs:
        if paragraph.text == exact:
            return paragraph
    raise ValueError(f"No se encontró el párrafo: {exact}")


def restart_numbering(doc, paragraphs, template) -> None:
    """Crea una lista real nueva que empieza en 1 usando el formato existente."""
    num_pr = template._p.pPr.numPr
    old_num_id = int(num_pr.numId.val)
    numbering = doc.part.numbering_part.element
    old_num = next(node for node in numbering.findall(qn("w:num"))
                   if int(node.get(qn("w:numId"))) == old_num_id)
    abstract_id = int(old_num.find(qn("w:abstractNumId")).get(qn("w:val")))
    new_num_id = max(int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))) + 1
    new_num = OxmlElement("w:num")
    new_num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    new_num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    new_num.append(override)
    numbering.append(new_num)
    for paragraph in paragraphs:
        p_pr = paragraph._p.get_or_add_pPr()
        current = p_pr.find(qn("w:numPr"))
        if current is not None:
            p_pr.remove(current)
        current = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), str(new_num_id))
        current.extend([ilvl, num_id])
        p_pr.append(current)


def remove_page_break_before(paragraph) -> None:
    previous = paragraph._p.getprevious()
    if previous is not None and previous.tag == qn("w:p") and previous.xpath(".//w:br[@w:type='page']"):
        previous.getparent().remove(previous)


def main() -> None:
    doc = Document(DOCX)

    replace_text(paragraph_by_text(doc, "Predicción diaria del precio spot, evaluación de modelos y optimización BESS"),
                 "Predicción diaria del precio spot, evaluación de modelos, optimización BESS y asistente del proyecto")
    replace_text(paragraph_by_text(doc, "Versión 1.0  |  3 de septiembre de 2026"),
                 "Versión 2.0  |  3 de septiembre de 2026")

    replace_text(paragraph_by_text(doc, "Pulso Energía es un dashboard de consulta para seguir el ciclo completo de una predicción de precio eléctrico: primero muestra el precio esperado por hora, después contrasta el comportamiento de los modelos y, por último, presenta el plan operativo de una batería cuando dicho plan ha sido calculado y guardado."),
                 "Pulso Energía es una aplicación de consulta para seguir el ciclo completo de una predicción de precio eléctrico: muestra el precio esperado por hora, contrasta el comportamiento de los modelos, presenta el plan operativo BESS guardado y permite consultar los datos y la metodología mediante un asistente contextual.")
    replace_text(paragraph_by_text(doc, "Usuarios y decisiones que soporta"), "Personas usuarias y decisiones que soporta")

    recorrido_bess = paragraph_by_text(doc, "BESS: traducir una predicción guardada en un plan de carga/descarga y consultar su resultado realizado.")
    insert_like(paragraph_by_text(doc, "Principio de trazabilidad"), recorrido_bess,
                "Asistente: consultar datos y metodología del proyecto en lenguaje natural, con trazabilidad de la herramienta utilizada.")

    shared = paragraph_by_text(doc, "En la versión publicada, el acceso a datos requiere la contraseña compartida del equipo.")
    replace_text(shared, "El acceso requiere una de las cinco cuentas individuales registradas para el equipo; todas tienen permisos funcionales de lectura.")
    target = paragraph_by_text(doc, "Comportamiento común de la interfaz")
    bullet_template = paragraph_by_text(doc, "Los procesos de producción escriben predicciones, precios reales, evaluaciones y planes BESS.")
    insert_like(target, bullet_template, "La sesión dura ocho horas y se identifica mediante una cookie HttpOnly y SameSite=Strict.")
    insert_like(target, bullet_template, "En HTTPS la cookie conserva Secure; cerrar sesión la elimina del navegador.")
    insert_like(target, bullet_template, "Las contraseñas, hashes y credenciales técnicas no se entregan al frontend ni se documentan en claro.")

    replace_text(paragraph_by_text(doc, "Tres modos independientes: Predicción, Evaluación y BESS."),
                 "Cuatro secciones independientes: Predicción, Evaluación, BESS y Asistente.")
    hours = paragraph_by_text(doc, "Las horas esperadas respetan días de 23, 24 o 25 horas según el cambio horario de Europe/Madrid.")
    replace_text(hours, "Las horas esperadas respetan días de 23, 24 o 25 horas según el cambio horario de Europe/Madrid y se presentan como H1, H2, …, H24 —o el número real de períodos del día—.")

    sources = paragraph_by_text(doc, "6. Fuentes de datos y reglas de presentación")
    insert_heading(sources, "6. Vista Asistente", 1)
    insert_like(sources, paragraph_by_text(doc, "Permite analizar el precio horario previsto para un día, comparar modelos y verificar qué ocurrió realmente cuando el mercado ya está cerrado."),
                "El asistente permite formular preguntas sobre precios, predicciones, batería, solar, sistema eléctrico y metodología del TFM. Utiliza el endpoint POST /api/asistente y la misma sesión de Pulso Energía; no tiene una contraseña independiente.")
    insert_heading(sources, "6.1 Orden de resolución", 2)
    numbered = paragraph_by_text(doc, "Predicción: entender el perfil horario esperado y compararlo con el precio real cuando esté disponible.")
    assistant_steps = [
        insert_like(sources, numbered, "Primero utiliza una herramienta determinista y previamente probada cuando la pregunta encaja."),
        insert_like(sources, numbered, "Para preguntas de metodología o decisiones, utiliza búsqueda semántica sobre la documentación del proyecto."),
        insert_like(sources, numbered, "Solo como último recurso utiliza SQL dinámico de lectura y lo advierte expresamente en la respuesta."),
    ]
    restart_numbering(doc, assistant_steps, numbered)
    insert_heading(sources, "6.2 Trazabilidad y disponibilidad", 2)
    insert_like(sources, bullet_template, "La respuesta identifica la fuente o herramienta empleada y puede incluir Markdown y gráficos.")
    insert_like(sources, bullet_template, "La clave de Anthropic tfm-equipo reside únicamente en el servidor y nunca se entrega al navegador.")
    insert_like(sources, bullet_template, "Si el asistente no está disponible, Predicción, Evaluación y BESS continúan funcionando.")

    replace_text(sources, "7. Fuentes de datos y reglas de presentación")
    dependencies_heading = paragraph_by_text(doc, "7. Dependencias operativas y controles")
    replace_text(dependencies_heading, "8. Dependencias operativas y controles")
    remove_page_break_before(dependencies_heading)
    glossary_heading = paragraph_by_text(doc, "8. Glosario funcional")
    replace_text(glossary_heading, "9. Glosario funcional")
    remove_page_break_before(glossary_heading)
    replace_text(paragraph_by_text(doc, "9. Mejora funcional pendiente"), "10. Mejora funcional pendiente")
    replace_text(paragraph_by_text(doc, "El punto abierto prioritario es cerrar el pipeline BESS: después de guardar las predicciones D+1 debe generarse y persistirse automáticamente un plan para un modelo operativo explícito. Actualmente no hay un modelo marcado como campeón y el fallback ensemble no está presente entre las predicciones más recientes; por ello la elección del modelo debe acordarse antes de automatizar el plan."),
                 "El punto abierto prioritario es cerrar el pipeline BESS: después de guardar las predicciones D+1 debe generarse y persistirse automáticamente un plan para un modelo operativo explícito. La elección del modelo debe acordarse antes de automatizar el proceso.")
    replace_text(paragraph_by_text(doc, "Decisión necesaria. Definir el modelo operativo que gobernará BESS (por ejemplo, GRU) y añadir el paso de planificación al proceso diario. Esta decisión no debe resolverse con un valor fijo oculto en la interfaz."),
                 "Decisión necesaria. Definir el modelo operativo que gobernará BESS y añadir la planificación al proceso diario; nunca resolverlo con un fallback oculto en la interfaz.",
                 bold_prefix="Decisión necesaria.")

    dependency_target = paragraph_by_text(doc, "Seguridad")
    insert_like(dependency_target, bullet_template,
                "Asistente: servicio interno del proyecto disponible detrás de Nginx, clave del proveedor configurada en servidor y sesión válida de Pulso Energía.")
    replace_text(paragraph_by_text(doc, "La sesión del equipo utiliza una cookie segura, HttpOnly y con caducidad de ocho horas."),
                 "Cada integrante utiliza su propia cuenta. La sesión usa una cookie HttpOnly, SameSite=Strict y con caducidad de ocho horas; en HTTPS también es Secure.")
    replace_text(paragraph_by_text(doc, "La instalación local se limita a 127.0.0.1 y puede ejecutarse sin contraseña de equipo."),
                 "La ejecución local se limita a 127.0.0.1. Cuando consulta el servidor, el proxy local conserva la sesión sin exponer la cookie ni las credenciales técnicas.")
    replace_text(paragraph_by_text(doc, "Cerrar sesión impide consultar datos hasta introducir nuevamente la contraseña del equipo."),
                 "Cerrar sesión impide consultar datos hasta introducir nuevamente el usuario y la contraseña de una cuenta válida.")
    controls_target = paragraph_by_text(doc, "Cerrar sesión impide consultar datos hasta introducir nuevamente el usuario y la contraseña de una cuenta válida.")
    insert_like(paragraph_by_text(doc, "8. Glosario funcional") if False else paragraph_by_text(doc, "9. Glosario funcional"),
                paragraph_by_text(doc, "Al abrir, la app selecciona el último día cerrado y permite avanzar al día futuro."),
                "El asistente reutiliza la sesión de la app, identifica la herramienta usada y falla de forma independiente del dashboard.")

    source_table = doc.tables[0]
    row = source_table.add_row()
    values = ("Asistente", "POST /api/asistente + herramientas", "Responde con datos o documentación y declara la fuente utilizada.")
    for cell, value in zip(row.cells, values):
        set_table_cell_text(cell, value)
        shade(cell, "F1F6F4")

    glossary = doc.tables[1]
    additions = [
        ("Herramienta determinista", "Función predefinida", "Código fijo y probado para una consulta concreta; es la vía preferente del asistente."),
        ("SQL dinámico", "Último recurso", "Consulta SELECT generada para una pregunta no cubierta; se identifica como menos verificada."),
    ]
    for values in additions:
        row = glossary.add_row()
        for cell, value in zip(row.cells, values):
            set_table_cell_text(cell, value)
            if len(glossary.rows) % 2 == 0:
                shade(cell, "F1F6F4")

    props = doc.core_properties
    props.title = "Pulso Energía — Definición funcional"
    props.subject = "Predicción, evaluación, BESS, acceso individual y asistente del proyecto"
    props.comments = "Versión 2.0 actualizada el 3 de septiembre de 2026"
    doc.save(DOCX)
    print(f"Actualizado: {DOCX}")
    print(f"Párrafos: {len(doc.paragraphs)} · Tablas: {len(doc.tables)}")


if __name__ == "__main__":
    main()
